from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / 'Kreston_NBB_Website_Strategy_and_Redesign_Plan.docx'
WEB_LOGO = ROOT / 'kreston-nbb-logo.png'
LOGO_PNG = ROOT / '.kreston_nbb_logo_for_report.png'

BLUE = '189CD8'
TEAL = '43BBC7'
SLATE = '243746'
INK = '0F172A'
MUTED = '6B7280'
PALE = 'EEF7FA'
PALE_ALT = 'F6F8FA'
BORDER = 'D9D9D9'
WHITE = 'FFFFFF'
BLACK = '000000'


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for margin, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{margin}'))
        if node is None:
            node = OxmlElement(f'w:{margin}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table, color=BORDER, size='6'):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = f'w:{edge}'
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement('w:cantSplit')
    cant_split.set(qn('w:val'), 'true')
    tr_pr.append(cant_split)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(int(inches * 1440)))
    tc_w.set(qn('w:type'), 'dxa')


def add_hyperlink(paragraph, text, url, color='0369A1'):
    part = paragraph.part
    rel_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), rel_id)
    new_run = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')
    r_color = OxmlElement('w:color')
    r_color.set(qn('w:val'), color)
    r_pr.append(r_color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement('w:t')
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def set_run_font(run, name='Arial', size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn('w:pBdr'))
    if p_bdr is None:
        p_bdr = OxmlElement('w:pBdr')
        p_pr.append(p_bdr)
    for edge in ('top', 'left', 'bottom', 'right', 'between', 'bar'):
        element = p_bdr.find(qn(f'w:{edge}'))
        if element is None:
            element = OxmlElement(f'w:{edge}')
            p_bdr.append(element)
        element.set(qn('w:val'), 'nil')


def add_page_number(paragraph, align_right=True):
    if align_right:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t')
    text.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    add_page_number(p, align_right=True)


def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph(style='Title')
    remove_paragraph_borders(p)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(text)
    set_run_font(r, size=28, bold=True, color=BLACK)
    if subtitle:
        s = doc.add_paragraph(style='Subtitle')
        remove_paragraph_borders(s)
        s.paragraph_format.space_after = Pt(18)
        sr = s.add_run(subtitle)
        set_run_font(sr, size=15, bold=False, color=BLACK)
        sr.italic = False
    return p


def add_heading(doc, text, level=1, page_break=False):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break
    r = p.add_run(text)
    set_run_font(r, bold=True, color=BLACK)
    return p


def add_para(doc, text='', bold_lead=None, style=None, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True, color=INK)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, color=INK)
    else:
        r = p.add_run(text)
        set_run_font(r, color=INK)
    return p


def add_bullets(doc, items, level=0):
    style = 'List Bullet' if level == 0 else 'List Bullet 2'
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, color=INK)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.27)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f'{index}.  {item}')
        set_run_font(r, color=INK)


def add_table(doc, headers, rows, widths=None, header_fill=SLATE, font_size=9.25):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = 'Table Grid'
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_fill(cell, header_fill)
        set_cell_margins(cell, top=110, bottom=110, start=110, end=110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(header)
        set_run_font(r, size=font_size, bold=True, color=WHITE)
        if widths:
            set_cell_width(cell, widths[idx])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for idx, value in enumerate(row):
            cell = cells[idx]
            if row_index % 2 == 1:
                set_cell_fill(cell, PALE_ALT)
            set_cell_margins(cell, top=95, bottom=95, start=110, end=110)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)
            if widths:
                set_cell_width(cell, widths[idx])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_source(doc, label, url):
    p = doc.add_paragraph(style='List Bullet')
    add_hyperlink(p, label, url)


def prepare_logo():
    if not WEB_LOGO.exists():
        return None
    im = Image.open(WEB_LOGO).convert('RGBA')
    bbox = im.getbbox()
    if bbox:
        im = im.crop(bbox)
    im.save(LOGO_PNG, 'PNG')
    return LOGO_PNG


def add_picture_with_alt(paragraph, image_path, width, alt):
    run = paragraph.add_run()
    shape = run.add_picture(str(image_path), width=width)
    doc_pr = shape._inline.docPr
    doc_pr.set('descr', alt)
    return shape


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    props = doc.core_properties
    props.title = 'Kreston NBB Saudi Website Strategy and Redesign Plan'
    props.subject = 'Website strategy information architecture user experience content and delivery requirements'
    props.keywords = 'Kreston NBB Saudi website strategy UX UI bilingual Arabic SEO accessibility'
    props.author = ''

    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    normal.font.size = Pt(10.75)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after in [
        ('Title', 28, 0, 12),
        ('Subtitle', 15, 0, 18),
        ('Heading 1', 18, 16, 8),
        ('Heading 2', 13.5, 12, 5),
        ('Heading 3', 11.5, 9, 4),
    ]:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style._element.rPr.rFonts.set(qn('w:ascii'), style.font.name)
        style._element.rPr.rFonts.set(qn('w:hAnsi'), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = name != 'Subtitle'
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style in ('List Bullet', 'List Bullet 2', 'List Number'):
        doc.styles[list_style].font.name = 'Arial'
        doc.styles[list_style].font.size = Pt(10.5)
        doc.styles[list_style].font.color.rgb = RGBColor.from_string(INK)

    add_footer(section)

    logo = prepare_logo()
    if logo:
        lp = doc.add_paragraph()
        lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lp.paragraph_format.space_after = Pt(28)
        add_picture_with_alt(lp, logo, Inches(2.45), 'Kreston NBB Saudi logo')

    eyebrow = doc.add_paragraph()
    eyebrow.paragraph_format.space_after = Pt(8)
    er = eyebrow.add_run('WEBSITE STRATEGY REPORT')
    set_run_font(er, size=9.5, bold=True, color=BLUE)
    add_title(
        doc,
        'Kreston NBB Saudi Website Strategy and Redesign Plan',
        'Structure Content User Experience and Delivery Requirements',
    )
    add_para(doc, 'Prepared 2 October 2026', style=None)
    add_para(doc, 'Purpose  Define a practical foundation for a client-facing demo and a production-ready website programme.', style=None)
    scope = doc.add_paragraph()
    scope.paragraph_format.space_before = Pt(28)
    sr = scope.add_run('Scope decision')
    set_run_font(sr, size=10.5, bold=True, color=BLACK)
    scope.add_run('\n')
    sr2 = scope.add_run('The website represents Kreston NBB Saudi. The wider Kreston network appears only as a supporting affiliation where it strengthens a relevant claim.')
    set_run_font(sr2, size=12, color=INK)
    doc.add_page_break()

    add_heading(doc, 'How to use this document', 1)
    add_para(doc, 'This report brings together the live-site audit, the proposed structure, the design system, page requirements and the decisions needed from Kreston NBB. It is detailed enough to guide wireframes and a polished demo. Sections that depend on client confirmation are identified so the demo can use provisional content without turning assumptions into published claims.')
    add_heading(doc, 'Contents', 2)
    contents = [
        'Executive recommendation',
        'Current website assessment',
        'Audience needs and user journeys',
        'Recommended website structure',
        'Homepage and page blueprints',
        'User experience and interface direction',
        'Content and messaging plan',
        'English and Arabic experience',
        'Search accessibility performance and privacy',
        'Content management and analytics',
        'Demo scope and delivery roadmap',
        'Client requirements and acceptance criteria',
        'Source and asset references',
    ]
    add_numbered(doc, contents)

    add_heading(doc, 'Executive recommendation', 1, page_break=True)
    add_para(doc, 'Kreston NBB should replace the current brochure-style website with a bilingual professional-services platform that helps a visitor answer four questions quickly: Is this firm credible in Saudi Arabia, does it offer the service I need, who will advise me, and how do I start a conversation? The present site supplies some useful brand and service material, but weak service pages, incomplete Arabic content and inconsistent contact information reduce trust at the point of decision.')
    add_para(doc, 'The new structure should centre on Kreston NBB itself. Local services, leadership, regulatory understanding and sector experience should carry the narrative. Kreston network membership can support cross-border work, but it should not occupy a top-level navigation item or a large share of the homepage.')
    add_heading(doc, 'Recommended first release', 2)
    add_bullets(doc, [
        'A focused homepage that routes visitors to services, experts, insights and consultation.',
        'A services hub and six complete service pages with clear outcomes and named contacts.',
        'An About section covering the Saudi firm, leadership, credentials and a short network affiliation statement.',
        'An Insights hub that combines articles, guides, news and events.',
        'Careers and Contact journeys with secure, privacy-aware forms.',
        'Complete English and Arabic experiences with true RTL layouts and human-reviewed content.',
    ])
    add_heading(doc, 'Working positioning', 2)
    add_para(doc, 'Clear advice for confident business decisions in Saudi Arabia.')
    add_para(doc, 'This is provisional copy for design and stakeholder discussion. The final line should be approved with the wider messaging and trademark guidance.')

    add_heading(doc, 'Current website assessment', 1)
    add_para(doc, 'The audit covered the English and Arabic homepages, About, Services, Kreston Global, Insights, News and Events, Careers, Contact and all six service-detail routes. The review also checked visible styles, assets, form structure, metadata, heading hierarchy, localisation and basic indexing files.')
    add_heading(doc, 'What is worth preserving', 2)
    add_bullets(doc, [
        'The Kreston NBB Saudi wordmark and blue-teal identity.',
        'The existing six service categories as a starting taxonomy.',
        'The Saudi market context and Riyadh presence.',
        'Kreston network membership as secondary evidence for international assignments.',
        'Insights as a useful way to demonstrate current knowledge when content is maintained.',
    ])
    add_heading(doc, 'Priority findings', 2)
    findings = [
        ('Critical', 'Service pages', 'All six detail routes render without useful service content.', 'Rebuild every service page around client needs, outcomes, sub-services, Saudi context, experts and a consultation CTA.'),
        ('Critical', 'Arabic experience', 'The Arabic homepage is partly translated and includes irrelevant industrial-contracting language.', 'Commission professional Arabic copy and test the complete RTL interface with native reviewers.'),
        ('Critical', 'Contact accuracy', 'The current footer and the official Kreston member listing show different address and phone details.', 'Confirm one canonical legal identity, address, phone, email and map pin before launch.'),
        ('Critical', 'Forms', 'Several forms use GET, unnamed fields, weak labels and no visible consent or validation.', 'Use secure server-side submission, proper labels, validation, consent, routing and success states.'),
        ('High', 'Positioning', 'The site uses generic claims and gives too much space to global network statistics.', 'Lead with NBB people, Saudi expertise, services and approved local credentials.'),
        ('High', 'Trust', 'Leadership, service leads, sector evidence, case work and client outcomes are absent.', 'Add named experts, approved credentials and evidence-backed experience.'),
        ('High', 'SEO', 'Descriptions are missing or placeholders, important pages lack H1s and no schema was detected.', 'Create unique metadata, semantic headings, structured data, descriptive URLs and internal links.'),
        ('High', 'Indexing', 'Robots and XML sitemap routes return 404.', 'Generate and monitor both files automatically.'),
        ('High', 'Freshness', 'The visible news content is dated and network statistics conflict.', 'Establish editorial ownership and review dates for claims and statistics.'),
        ('Medium', 'Accessibility', 'Heading levels skip, generic alt text is common and some actions lack accessible names.', 'Build and test to WCAG 2.2 AA with keyboard, screen-reader, contrast and reduced-motion checks.'),
        ('Medium', 'Visual consistency', 'Inter and Tahoma are mixed without a clear system; some interaction text is only 12 px.', 'Use one English type system, one approved Arabic companion and accessible component specifications.'),
    ]
    add_table(doc, ['Priority', 'Area', 'Observed issue', 'Recommended response'], findings, widths=[0.72, 1.05, 2.27, 3.02], font_size=8.5)

    add_heading(doc, 'Specific defects to correct', 2)
    add_bullets(doc, [
        'The footer phone link uses the invalid phone scheme and appears to carry an email value. Use a verified tel link.',
        'A social link labelled YouTube points to X or Twitter. Match every label, icon and destination.',
        'The Contact page exposes address and email headings without complete visible values.',
        'The current About copy contains awkward or inaccurate wording, including Forum of Forms.',
        'Service, Career and Contact pages lack a strong semantic H1 structure.',
        'Insight URLs use database identifiers rather than readable slugs.',
        'The autoplay hero video needs a poster, reduced-motion behaviour and a lighter mobile alternative.',
    ])

    add_heading(doc, 'Audience needs and user journeys', 1)
    add_para(doc, 'The site should organise information around the decisions visitors make, not around the internal structure of the firm. The following audience groups are appropriate for the demo and should be confirmed before production.')
    audience_rows = [
        ('Saudi business owner or finance leader', 'Find reliable audit, tax, accounting or advisory support.', 'Relevant service scope, Saudi knowledge, senior contact, clear next step.'),
        ('Regional or international company entering Saudi Arabia', 'Understand local obligations and find an adviser who can coordinate across borders.', 'Saudi setup and compliance context, cross-border capability, named lead.'),
        ('Regulated or complex organisation', 'Verify credentials, governance and specialist capability before engaging.', 'Approved licences, CMA wording, sector experience, quality controls.'),
        ('Prospective employee', 'Understand the firm, opportunities and application process.', 'Culture, roles, expectations, privacy-aware application route.'),
        ('Existing client or referral partner', 'Find a contact, article or related service quickly.', 'Search, adviser profiles, direct contact and current insights.'),
    ]
    add_table(doc, ['Audience', 'Primary task', 'Evidence needed'], audience_rows, widths=[1.6, 2.35, 3.1], font_size=9)

    add_heading(doc, 'Priority user journeys', 2)
    add_heading(doc, 'Service enquiry', 3)
    add_numbered(doc, [
        'Arrive through search, referral or the homepage.',
        'Recognise the relevant service and Saudi context within the first screen.',
        'Review outcomes, sub-services, approach, credentials and the lead adviser.',
        'Choose Request a consultation and select the relevant service.',
        'Receive a clear confirmation and realistic response expectation.',
    ])
    add_heading(doc, 'Credibility check', 3)
    add_numbered(doc, [
        'Open About, Leadership or a service page.',
        'Verify the legal firm, local credentials and leadership.',
        'Understand how network affiliation supports rather than replaces the local engagement.',
        'Move directly to an adviser profile or consultation route.',
    ])
    add_heading(doc, 'Insight to enquiry', 3)
    add_numbered(doc, [
        'Land on a Saudi-specific article from search or social media.',
        'See the author, date, expertise and review status.',
        'Follow a relevant service or related article.',
        'Contact the responsible adviser without returning to the homepage.',
    ])

    add_heading(doc, 'Recommended website structure', 1)
    add_para(doc, 'The primary navigation should stay short. Kreston Global should not appear as a top-level item. Its existing page can redirect to a short affiliation section within About unless the client has a contractual reason to keep a separate page.')
    add_heading(doc, 'Primary navigation', 2)
    nav_rows = [
        ('Home', 'Firm proposition, services, trust, expertise and consultation.'),
        ('Services', 'Six service categories and detailed service pages.'),
        ('Industries', 'Only sectors Kreston NBB can support with evidence and named expertise.'),
        ('Insights', 'Articles, guides, news and events in one maintained hub.'),
        ('About', 'Firm, leadership, credentials, quality and network affiliation.'),
        ('Careers', 'Employer proposition, roles and application process.'),
        ('Contact', 'Verified office details and routed enquiry form.'),
    ]
    add_table(doc, ['Navigation item', 'Purpose'], nav_rows, widths=[1.4, 5.65], font_size=9.5)
    add_para(doc, 'Utility actions should include the English and Arabic switch and one primary Request a consultation button. Site search can be introduced when the insight library is large enough to justify it; it does not need to occupy primary space in the demo.')

    add_heading(doc, 'Proposed sitemap', 2)
    sitemap_rows = [
        ('Home', 'Home'),
        ('Services', 'Hub; Audit and Assurance; Internal Audit Risk and Compliance; Tax and Zakat; Accounting and Advisory; Management Consulting; Operations and Technology'),
        ('Industries', 'Hub plus confirmed industry pages'),
        ('Insights', 'Hub; Article; Guide; News; Event'),
        ('About', 'Our Firm; Leadership and Team; Credentials and Quality; Network Affiliation'),
        ('Careers', 'Careers plus role and application route'),
        ('Contact', 'Contact'),
        ('Legal', 'Privacy; Cookies; Terms; Accessibility'),
    ]
    add_table(doc, ['Section', 'Pages'], sitemap_rows, widths=[1.3, 5.75], font_size=9)
    add_para(doc, 'Industries should remain provisional until NBB confirms priority sectors, evidence and responsible advisers. The current external member profile lists many sectors, but publishing a long generic list would weaken credibility.')

    add_heading(doc, 'URL and redirect approach', 2)
    add_bullets(doc, [
        'Use readable service and article slugs instead of database IDs.',
        'Keep one predictable language pattern, such as /en and /ar.',
        'Create one-to-one permanent redirects from every current indexable URL.',
        'Redirect /kreston-global to the NBB About affiliation section unless it must remain separate.',
        'Retain query parameters only where they support deliberate filtering or campaign attribution.',
    ])

    add_heading(doc, 'Page purpose and conversion plan', 2)
    page_rows = [
        ('Home', 'Explain why NBB is relevant and route the visitor.', 'Request a consultation'),
        ('Services hub', 'Help visitors identify the right capability.', 'View a service'),
        ('Service detail', 'Demonstrate scope, expertise and outcomes.', 'Speak to the service lead'),
        ('Industries hub', 'Show where NBB has relevant experience.', 'Explore sector expertise'),
        ('Industry detail', 'Connect sector issues to services and experts.', 'Discuss a sector need'),
        ('About', 'Establish the firm, leadership and credentials.', 'Meet the team'),
        ('Leadership profile', 'Make expertise and contact personal.', 'Contact this adviser'),
        ('Insights hub', 'Surface current, useful Saudi guidance.', 'Read an insight'),
        ('Article', 'Demonstrate expertise and support search discovery.', 'Discuss the issue'),
        ('Careers', 'Explain the employee proposition and process.', 'View roles or apply'),
        ('Contact', 'Reduce uncertainty and route the enquiry.', 'Submit enquiry'),
    ]
    add_table(doc, ['Page', 'Primary purpose', 'Primary action'], page_rows, widths=[1.25, 3.75, 2.05], font_size=9)

    add_heading(doc, 'Homepage blueprint', 1)
    add_para(doc, 'The homepage should help a first-time visitor understand the firm and reach the right next step without reading a corporate history. Each section below has a distinct job in that journey.')
    home_rows = [
        ('1', 'Header', 'Logo, concise navigation, language switch and consultation button.', 'Persistent orientation and action.'),
        ('2', 'Hero', 'NBB-focused value proposition, short supporting copy, primary and secondary CTA, authentic Saudi or team image.', 'Immediate relevance.'),
        ('3', 'Trust strip', 'Approved Saudi credentials, CMA wording if approved, Kreston membership and one local proof point.', 'Early reassurance.'),
        ('4', 'Services', 'Six outcome-led cards with short, distinct descriptions.', 'Service discovery.'),
        ('5', 'Why NBB', 'Saudi regulatory understanding, senior attention and responsive delivery.', 'Clear differentiation.'),
        ('6', 'Industries', 'Four to six evidence-backed priority sectors.', 'Audience recognition.'),
        ('7', 'Leadership', 'Managing partner and selected service leads with direct profile links.', 'Human trust.'),
        ('8', 'Featured insight', 'One current Saudi business, tax, audit or regulatory article.', 'Demonstrated expertise.'),
        ('9', 'Connected when needed', 'One concise cross-border capability statement.', 'Affiliation support without brand drift.'),
        ('10', 'Consultation band', 'Short invitation and one primary action.', 'Conversion.'),
        ('11', 'Footer', 'Verified details, legal pages, social links and member-firm disclaimer.', 'Closure and compliance.'),
    ]
    add_table(doc, ['Order', 'Section', 'Content', 'User outcome'], home_rows, widths=[0.45, 1.15, 3.8, 1.65], font_size=8.5)
    add_heading(doc, 'Provisional hero copy', 2)
    add_para(doc, 'Saudi insight Confident decisions')
    add_para(doc, 'Audit, tax and advisory expertise for organisations navigating growth, regulation and change in Saudi Arabia.')
    add_bullets(doc, [
        'Primary action  Request a consultation',
        'Secondary action  Explore our services',
    ])
    add_para(doc, 'The final hero must state what NBB does and where it operates. It should not open with network scale, abstract business language or an unsupported best-in-market claim.')

    add_heading(doc, 'Page blueprints', 1)
    add_heading(doc, 'Services hub', 2)
    add_bullets(doc, [
        'Introductory statement explaining how NBB supports organisations in Saudi Arabia.',
        'Six service cards with distinct outcomes rather than interchangeable descriptions.',
        'Short method or engagement approach.',
        'Relevant industries and insights.',
        'Consultation CTA with service preselected.',
    ])
    add_heading(doc, 'Service detail template', 2)
    add_numbered(doc, [
        'Service-specific hero with outcome, audience and named lead.',
        'Business problems the service addresses.',
        'Sub-services and expected deliverables.',
        'Saudi regulatory and market context.',
        'How an engagement works.',
        'Relevant sectors and approved experience.',
        'Lead adviser profile and contact route.',
        'Related insight and useful FAQs.',
        'Consultation CTA and clear privacy statement.',
    ])
    add_heading(doc, 'About', 2)
    add_bullets(doc, [
        'Accurate legal identity and concise local firm story.',
        'Leadership and team with roles, biographies and approved professional links.',
        'Credentials, licences, governance and quality controls.',
        'Values explained through client behaviour rather than generic adjectives.',
        'Short network affiliation statement and approved independent-member disclaimer.',
        'Riyadh presence and verified contact details.',
    ])
    add_heading(doc, 'Industries', 2)
    add_para(doc, 'Industry pages should launch only when NBB can name the client problems, relevant services and responsible experts for that sector. Each page should address current Saudi context, common issues, service connections, proof and a direct consultation route.')
    add_heading(doc, 'Insights', 2)
    add_bullets(doc, [
        'One hub for articles, guides, news and events.',
        'Useful filters by topic, service and industry.',
        'Article date, author, role, review date and reading time.',
        'Related services, related content and author contact.',
        'Content owner and review or expiry process.',
    ])
    add_heading(doc, 'Careers', 2)
    add_bullets(doc, [
        'Firm culture and employee proposition supported by real team content.',
        'Current roles or a clear speculative-application policy.',
        'Simple application with only necessary personal information.',
        'Recruitment privacy notice, retention period and secure file upload.',
        'Remove date of birth and gender unless the client establishes a necessary and lawful basis.',
    ])
    add_heading(doc, 'Contact', 2)
    add_bullets(doc, [
        'Verified legal office address, map pin, phone, email and working hours.',
        'Topic and service selector for routing.',
        'Short form with name, work email, phone if optional, organisation, topic and message.',
        'Clear consent, error, success and response-expectation messages.',
        'Separate recruitment and general enquiry routes.',
    ])

    add_heading(doc, 'User experience and interface direction', 1)
    add_para(doc, 'The design should feel calm, precise and current. The strongest improvement will come from hierarchy, spacing, real expertise and predictable interaction rather than decorative effects.')
    add_heading(doc, 'Brand palette', 2)
    palette_rows = [
        ('Kreston blue', '#189CD8', 'Brand accent, selected graphics and focus treatment.'),
        ('Kreston teal', '#43BBC7', 'Secondary brand accent and gradient end.'),
        ('Deep slate', '#243746', 'Headings, dark sections and strong contrast.'),
        ('Ink', '#0F172A', 'Primary body text.'),
        ('Canvas', '#EEEFF4', 'Alternating section background.'),
        ('White', '#FFFFFF', 'Primary surface and card background.'),
        ('Muted grey', '#6B7280', 'Supporting text when contrast remains sufficient.'),
        ('Border grey', '#E5E7EB', 'Inputs, dividers and card edges.'),
    ]
    table = add_table(doc, ['Colour', 'Value', 'Recommended role'], palette_rows, widths=[1.4, 1.0, 4.65], font_size=9.25)
    for idx, (_, value, _) in enumerate(palette_rows, start=1):
        set_cell_fill(table.cell(idx, 1), value.replace('#', ''))
        r = table.cell(idx, 1).paragraphs[0].runs[0]
        if value in ('#243746', '#0F172A', '#6B7280'):
            r.font.color.rgb = RGBColor.from_string(WHITE)
        else:
            r.font.color.rgb = RGBColor.from_string(BLACK)

    add_heading(doc, 'Colour accessibility', 2)
    add_para(doc, 'The bright brand blue and teal should not carry small white text because the contrast is insufficient. Use deep slate or an accessible darker blue for text buttons, and reserve the bright colours for larger graphics, focus outlines, icons and non-text accents. Every final state must be checked rather than inferred from the palette.')
    add_heading(doc, 'Typography', 2)
    add_bullets(doc, [
        'English  Inter in weights 400, 500, 600 and 700.',
        'Arabic  Noto Sans Arabic or another approved companion; Tahoma only as fallback.',
        'Body copy  16 to 18 px on the website with a 1.5 to 1.7 line height.',
        'Metadata  no smaller than 14 px.',
        'Maximum reading measure  approximately 65 to 75 characters per line.',
    ])
    add_heading(doc, 'Layout and spacing', 2)
    add_bullets(doc, [
        'Use a responsive 12-column desktop grid and a single-column mobile reading flow.',
        'Keep generous outer gutters and consistent vertical section spacing.',
        'Avoid dense card walls; alternate short editorial sections with clear visual anchors.',
        'Use one main action per section and maintain the same label for the same action.',
        'Allow Arabic text more vertical and horizontal room because line breaks differ from English.',
    ])
    add_heading(doc, 'Photography and graphic language', 2)
    add_bullets(doc, [
        'Prioritise authentic leadership, team and work-context photography.',
        'Use Riyadh context selectively; avoid relying on skyline footage as the whole identity.',
        'Use a single line-icon family and the Kreston chevron as a subtle crop or directional device.',
        'Avoid stock handshakes, generic glass towers and unapproved AI-generated people.',
        'Request originals and usage rights for every migrated image and video.',
    ])
    add_heading(doc, 'Interaction patterns', 2)
    add_bullets(doc, [
        'Sticky header that reduces cleanly after scroll without hiding orientation.',
        'Visible hover and keyboard-focus states.',
        'Cards that expose a clear text link rather than relying on the entire card area.',
        'Accordion only for secondary FAQs, not for essential service content.',
        'Inline validation with a summary for form errors.',
        'Motion limited to short, functional transitions and disabled when reduced motion is requested.',
    ])
    add_heading(doc, 'Core component library', 2)
    component_rows = [
        ('Navigation', 'Desktop header, mobile menu, breadcrumb, language switch.'),
        ('Actions', 'Primary button, secondary button, text link, adviser contact.'),
        ('Content', 'Service card, industry card, article card, person card, credential list.'),
        ('Proof', 'Statistic, licence or membership marker, testimonial, case result.'),
        ('Editorial', 'Article header, author line, related content, download pattern.'),
        ('Forms', 'Field, select, textarea, consent, upload, error and success states.'),
        ('Navigation aids', 'Pagination, filters, search results and back links.'),
        ('System', 'Cookie controls, 404, empty state, loading state and maintenance message.'),
    ]
    add_table(doc, ['Component group', 'Required patterns'], component_rows, widths=[1.55, 5.5], font_size=9.25)

    add_heading(doc, 'Mobile experience', 2)
    add_bullets(doc, [
        'Place the value proposition and primary action before the hero media.',
        'Replace autoplay video with an optimised still image or user-initiated media.',
        'Keep tap targets at least 44 by 44 CSS pixels.',
        'Collapse navigation by task and keep the language switch easy to reach.',
        'Avoid horizontal carousels for critical content.',
        'Use phone and map links only after the details are verified.',
    ])

    add_heading(doc, 'Content and messaging plan', 1)
    add_para(doc, 'The writing should sound like experienced advisers speaking clearly to a business decision-maker. It should make specific claims that NBB can support, state who a service helps and show the next action.')
    add_heading(doc, 'Message hierarchy', 2)
    message_rows = [
        ('Firm promise', 'Clear advice for confident business decisions in Saudi Arabia.'),
        ('Local value', 'Guidance grounded in Saudi regulation and business realities.'),
        ('Relationship value', 'Direct access to experienced professionals who understand the client.'),
        ('Connected capability', 'Access to relevant Kreston specialists when an engagement crosses borders.'),
        ('Action', 'Request a consultation with the appropriate NBB adviser.'),
    ]
    add_table(doc, ['Message level', 'Working expression'], message_rows, widths=[1.55, 5.5], font_size=9.5)
    add_heading(doc, 'Tone rules', 2)
    add_bullets(doc, [
        'Use precise, direct language and short paragraphs.',
        'Explain services through client problems, work performed and outcomes.',
        'Replace generic superlatives with credentials, examples or responsible experts.',
        'Keep necessary legal and regulatory qualifications close to the claim.',
        'Write Arabic as native professional copy rather than a literal translation of English.',
    ])
    add_heading(doc, 'Claims to avoid unless evidenced', 2)
    add_bullets(doc, [
        'Top accounting firm',
        'Best financial services',
        'Top-ranked tax services',
        'Leading or largest without a named source, scope and date',
        'Guaranteed results or absolute compliance claims',
    ])
    add_heading(doc, 'Service naming', 2)
    add_para(doc, 'The final labels should use terminology that Saudi clients recognise and that NBB can substantiate. Audit and Assurance, Internal Audit Risk and Compliance, Tax and Zakat, Accounting and Advisory, Management Consulting, and Operations and Technology are suitable working labels. Client review should confirm whether Financial Audit, Forensic, Risk, VAT, Transfer Pricing or other specialist terms deserve separate visibility.')

    add_heading(doc, 'English and Arabic experience', 1)
    add_para(doc, 'English and Arabic should be treated as two complete editorial experiences that share one content model. The Arabic site cannot be a partially translated layer over English pages.')
    add_heading(doc, 'Required bilingual behaviour', 2)
    add_bullets(doc, [
        'Preserve the equivalent page when the visitor changes language.',
        'Set language and direction attributes correctly on every page.',
        'Mirror layout and directional icons where meaning changes in RTL.',
        'Keep logos, photographs and non-directional brand marks unchanged.',
        'Use Arabic labels, validation messages, metadata, alt text and structured data.',
        'Allow independent publishing states so an incomplete translation cannot appear by accident.',
        'Support English and Arabic search, slugs and filtering.',
    ])
    add_heading(doc, 'Arabic quality checks', 2)
    add_bullets(doc, [
        'Native-language editorial review for tone and terminology.',
        'Service and regulatory terminology approved by NBB specialists.',
        'Visual review at mobile and desktop widths.',
        'Keyboard and screen-reader checks in RTL.',
        'No English fallback text in navigation, cards, forms, errors or footer.',
    ])

    add_heading(doc, 'Search accessibility performance and privacy', 1)
    add_heading(doc, 'Search optimisation', 2)
    add_bullets(doc, [
        'Unique title, description, canonical URL and H1 for every indexable page.',
        'Descriptive service and article slugs.',
        'XML sitemap, robots file, clean status codes and redirect monitoring.',
        'English and Arabic hreflang pairs, including self-references.',
        'Organisation or professional-service, service, person, article and breadcrumb structured data where supported.',
        'Consistent legal name, address and phone across the website and trusted external profiles.',
        'Open Graph and social images for important pages.',
        'Editorial internal links between services, industries, experts and insights.',
    ])
    add_heading(doc, 'Accessibility standard', 2)
    add_para(doc, 'Use WCAG 2.2 AA as the design and testing target. Automated checks are useful, but keyboard, screen-reader, zoom, contrast and RTL reviews must be performed by people.')
    add_bullets(doc, [
        'Logical heading hierarchy and landmarks.',
        'Visible focus and complete keyboard access.',
        'Text and control contrast that passes in every state.',
        'Persistent field labels, helpful instructions and error recovery.',
        'Meaningful alt text for informative images and empty alt text for decorative images.',
        'Captions or transcripts for meaningful video.',
        'Reduced-motion behaviour and no motion-dependent information.',
        'Usable content at 200 percent zoom without horizontal reading scroll.',
    ])
    add_heading(doc, 'Performance targets', 2)
    add_bullets(doc, [
        'Aim for Largest Contentful Paint below 2.5 seconds at the 75th percentile.',
        'Aim for Interaction to Next Paint below 200 milliseconds at the 75th percentile.',
        'Aim for Cumulative Layout Shift below 0.1 at the 75th percentile.',
        'Optimise responsive images, fonts and third-party scripts.',
        'Provide a poster and lightweight mobile alternative for hero media.',
        'Load non-critical media and embeds only when needed.',
    ])
    add_heading(doc, 'Forms privacy and security', 2)
    add_bullets(doc, [
        'Submit through secure server-side endpoints, not query strings.',
        'Use rate limiting, spam protection and safe validation.',
        'Collect only fields needed for the stated purpose.',
        'Show a clear privacy notice and retention basis at the form.',
        'Scan uploaded CVs, restrict file type and size, and limit access.',
        'Route enquiries to accountable owners and avoid exposing personal inboxes in code.',
        'Log success and failure without storing unnecessary message content in analytics.',
    ])

    add_heading(doc, 'Content management and analytics', 1)
    add_heading(doc, 'Required content types', 2)
    content_rows = [
        ('Service', 'Name, summary, audience, problems, sub-services, Saudi context, approach, lead, FAQs, CTA, SEO.'),
        ('Industry', 'Name, issues, services, experts, proof, insights, CTA, SEO.'),
        ('Person', 'Name, role, biography, services, industries, credentials, photo, approved contact links.'),
        ('Insight', 'Title, summary, body, author, publish date, review date, topics, related services, social image, SEO.'),
        ('Credential', 'Name, issuer, approved wording, evidence, effective date, review date.'),
        ('Office', 'Legal name, address, map, phone, email, hours and language variants.'),
        ('Site settings', 'Navigation, footer, social links, legal links, default metadata and form routing.'),
    ]
    add_table(doc, ['Content type', 'Minimum fields'], content_rows, widths=[1.3, 5.75], font_size=8.9)
    add_heading(doc, 'Publishing controls', 2)
    add_bullets(doc, [
        'Draft, review, approve, schedule and archive states.',
        'Separate English and Arabic status with missing-translation warnings.',
        'Named owner and review date for credentials, statistics and regulatory content.',
        'Role-based access for editors, reviewers and administrators.',
        'Preview for desktop, mobile, English and Arabic before publication.',
        'Redirect management and an audit trail for material content changes.',
    ])
    add_heading(doc, 'Measurement plan', 2)
    add_para(doc, 'Analytics should measure useful intent rather than page views alone. The implementation must follow the approved consent and privacy approach.')
    analytics_rows = [
        ('Consultation CTA', 'consultation_cta', 'Location, service and language.'),
        ('Form start', 'form_start', 'Form type and language.'),
        ('Successful enquiry', 'form_submit_success', 'Form type and routed service; no message text.'),
        ('Service engagement', 'service_view', 'Service and source page.'),
        ('Adviser contact', 'adviser_contact', 'Adviser and contact method.'),
        ('Language change', 'language_switch', 'Source and target language.'),
        ('Content download', 'content_download', 'Asset and topic.'),
        ('Career application', 'career_apply', 'Role or speculative route; no applicant details.'),
    ]
    add_table(doc, ['Interaction', 'Suggested event', 'Useful context'], analytics_rows, widths=[1.65, 1.75, 3.65], font_size=8.9)

    add_heading(doc, 'Demo scope and delivery roadmap', 1)
    add_heading(doc, 'Recommended demo', 2)
    add_para(doc, 'The demo should prove the navigation, visual system, service storytelling, conversion path and bilingual layout. It does not need every production integration or every article.')
    add_bullets(doc, [
        'Responsive homepage.',
        'Services hub.',
        'One complete service-detail example.',
        'About page with provisional leadership layout.',
        'Contact page and form states.',
        'One insight card and article template.',
        'English pages and representative Arabic RTL screens.',
        'Desktop and mobile responsive behaviour.',
    ])
    add_heading(doc, 'Demo assumptions', 2)
    add_bullets(doc, [
        'Extracted colours and logo are provisional until official brand assets arrive.',
        'Copy is working copy for structure and design, not approved publication copy.',
        'Leadership photos and biographies may use clearly marked placeholders.',
        'Credentials and contact details must not be invented; unconfirmed items remain labelled for approval.',
        'Forms can demonstrate interface states without sending data to a live mailbox.',
        'Network references remain brief and secondary.',
    ])
    add_heading(doc, 'Delivery phases', 2)
    phase_rows = [
        ('1 Discovery confirmation', 'Confirm audiences, services, legal identity, credentials, brand rules, languages and success measures.'),
        ('2 Content and structure', 'Approve sitemap, content model, page briefs, redirect map and bilingual workflow.'),
        ('3 UX and visual design', 'Produce wireframes, component system, desktop and mobile designs, and RTL variants.'),
        ('4 Demo build', 'Build representative pages and test navigation, responsiveness, forms and language behaviour.'),
        ('5 Production build', 'Connect the CMS, complete all content, integrations, SEO, analytics and security controls.'),
        ('6 Quality assurance', 'Test content, browsers, devices, accessibility, performance, SEO, forms, redirects and analytics.'),
        ('7 Launch and improvement', 'Monitor errors, enquiries, search performance and content freshness; schedule regular reviews.'),
    ]
    add_table(doc, ['Phase', 'Output'], phase_rows, widths=[1.75, 5.3], font_size=9.2)

    add_heading(doc, 'Client requirements', 1)
    add_para(doc, 'The following decisions and materials are needed before the demo becomes a production specification.')
    checklist_groups = [
        ('Brand', [
            'Official Kreston NBB logo files and usage rules.',
            'Applicable Kreston network brand requirements.',
            'Approved English and Arabic typefaces.',
            'Original photography and video with usage rights.',
        ]),
        ('Business and compliance', [
            'Exact legal entity name in English and Arabic.',
            'Approved wording and scope for licences and CMA status.',
            'Member-firm and liability disclaimer.',
            'Canonical office address, map pin, phone, email and working hours.',
            'Privacy, cookies, terms, recruitment-retention and accessibility policies.',
            'Hosting, data-residency and security requirements.',
        ]),
        ('Audience and conversion', [
            'Priority client types, company sizes and decision-makers.',
            'Priority industries and cross-border needs.',
            'The main reasons clients choose NBB.',
            'Lead owners, routing rules and response commitment.',
            'CRM, email, analytics and marketing integrations.',
        ]),
        ('Content', [
            'Approved descriptions for every service and sub-service.',
            'Leadership biographies, roles, photos and approved professional links.',
            'Credentials, awards, memberships and effective dates.',
            'Approved client logos, testimonials and case studies.',
            'Content migration and retirement decisions.',
            'English and Arabic editors and approvers.',
        ]),
        ('Delivery', [
            'CMS preference and internal editor roles.',
            'Domain, DNS, hosting and deployment ownership.',
            'Accessibility and browser support commitments.',
            'Budget, target launch date and approval milestones.',
            'Post-launch maintenance owner and service level.',
        ]),
    ]
    for heading, items in checklist_groups:
        add_heading(doc, heading, 2)
        add_bullets(doc, [f'Confirm  {item}' for item in items])

    add_heading(doc, 'Acceptance criteria', 1)
    add_para(doc, 'The production website should not launch until the following conditions are met.')
    acceptance = [
        'Every service page contains approved, useful content and a named conversion path.',
        'English and Arabic are complete, equivalent and professionally reviewed.',
        'Contact details, credentials and legal identity match approved records.',
        'Kreston network references remain accurate, approved and secondary to NBB.',
        'Every indexable page has a unique title, description, canonical and H1.',
        'Sitemap, robots, hreflang, structured data and redirects are verified.',
        'Forms are secure, labelled, validated, privacy-aware and tested end to end.',
        'WCAG 2.2 AA checks cover keyboard, focus, contrast, labels, zoom, screen readers and reduced motion.',
        'Mobile, tablet, desktop and RTL layouts are visually verified.',
        'Performance meets agreed Core Web Vitals targets on representative devices.',
        'Analytics records the agreed events without collecting sensitive form content.',
        'Every dated claim has an owner and review date.',
        'A post-launch owner is responsible for insights, security updates and ongoing quality.',
    ]
    add_bullets(doc, acceptance)

    add_heading(doc, 'Source and asset references', 1)
    add_para(doc, 'The report uses the public website and official Kreston references to identify current content, contradictions and opportunities. Publication claims still require Kreston NBB approval.')
    add_heading(doc, 'Web sources', 2)
    add_source(doc, 'Kreston NBB Saudi website', 'https://kreston-nbb.com/')
    add_source(doc, 'Kreston NBB services page', 'https://kreston-nbb.com/services')
    add_source(doc, 'Kreston NBB about page', 'https://kreston-nbb.com/about')
    add_source(doc, 'Kreston NBB official member profile', 'https://www.kreston.com/members/kreston-nbb/')
    add_source(doc, 'Kreston Global network overview used only to verify current network references', 'https://www.kreston.com/about-kreston-global/serving-your-international-business-needs/')
    add_heading(doc, 'Current asset inventory', 2)
    asset_rows = [
        ('Logo', '/images/logo/logo3.png', 'Request SVG master and reversed version.'),
        ('Hero video', '/uploads/upload1693739760796.mp4', 'Review rights, weight, poster and mobile fallback.'),
        ('Service icons', '/uploads/*.svg', 'Keep only if the style and rights are approved.'),
        ('Global map', '/images/mapwhite.png', 'Not required for the NBB-focused redesign.'),
        ('Network icons', '/images/icon/*.png', 'Do not migrate automatically; network statistics are secondary.'),
        ('About and insight media', '/uploads/*', 'Request originals, captions, alt text and rights.'),
        ('Fonts', 'Inter and Tahoma files', 'Confirm licensed production use and the Arabic companion font.'),
    ]
    add_table(doc, ['Asset', 'Current path', 'Action'], asset_rows, widths=[1.35, 2.2, 3.5], font_size=8.9)
    add_heading(doc, 'Implementation artefacts already prepared', 2)
    add_bullets(doc, [
        'Website discovery brief with the current audit and requirements checklist.',
        'Provisional CSS design tokens based on the extracted palette.',
        'NBB-first sitemap, homepage structure and page requirements in this report.',
    ])

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == '__main__':
    print(build_document())
